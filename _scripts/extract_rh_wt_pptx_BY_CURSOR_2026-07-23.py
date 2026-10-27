"""
Extract RH Window Treatments PPT slides → JSON for Builder WO import.
Skips cover/legend and HOLD/NO slides.
"""
import json
import re
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PPT = r"C:\Users\cindy\CCH Design\Cloud - Rolling Hills - Documents\Working Documents\Layout Custom Work orders\RH - Window Treatments Rev 2 8.4.2025.pptx"
OUT = r"C:\dev\CCH-Platform-Deploy\cch-deploy\_scripts\_tmp_rh_wt_slides_BY_CURSOR.json"

def walk(shape):
    lines = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            lines.extend(walk(s))
        return lines
    if shape.has_text_frame:
        t = shape.text_frame.text.strip()
        if t:
            lines.append(t)
    return lines

def clean(s):
    return re.sub(r"\s+", " ", (s or "").replace("\u00a0", " ")).strip()

def infer_category(blob):
    # Every slide carries "Box to hide roller shade" boilerplate — strip before classify
    b = blob
    for phrase in (
        "Box to hide roller shade. Drywall",
        "Box to hide roller shade",
        "Window Treatment Design",
        "Fabrics",
        "Design Specifications",
    ):
        b = re.sub(re.escape(phrase), " ", b, flags=re.I)
    b = b.lower()
    if "valance only" in b or "decorative valance" in b or "valance over tub" in b:
        return "drapery"
    if "roman shade" in b or re.search(r"\broman\b", b):
        return "roman"
    if "draper" in b or "drapery" in b:
        # drapery + shade on same slide (e.g. W5) → still roller primary if solar/blackout roller dominant
        if ("solar shade" in b or "roller solar" in b or "black out privacy shade" in b or "dual shade" in b) and "sheer draper" not in b:
            if "sheer draper" in b or "operable draper" in b:
                return "roller"  # layered; primary shade WO
            return "roller"
        return "drapery"
    if "dual shade" in b or "solar shade" in b or "roller solar" in b or "black out privacy shade" in b or "roller shade" in b:
        return "roller"
    return "roller"

def parse_header(texts):
    tag = ""
    room = ""
    header = ""
    for t in texts:
        if re.search(r"work\s*order", t, re.I) or t.lower().startswith("ork order"):
            header = t
            m = re.search(r"(?:WINDOW|DRAPERY)\s+TREATMENT\s*(.*)$", t, re.I | re.S)
            if m:
                tag = clean(m.group(1))
        if re.search(r"^Location:", t, re.I):
            room = clean(re.sub(r"^Location:\s*", "", t, flags=re.I))
        if "Location:" in t and "Work Order" in t:
            # combined box
            parts = t.split("Location:", 1)
            if len(parts) == 2:
                room = clean(parts[1].split("\n")[0])
                hm = re.search(r"(?:WINDOW|DRAPERY)\s+TREATMENT\s*(.*?)(?:\n|Location:|$)", t, re.I | re.S)
                if hm:
                    tag = clean(hm.group(1))
    return tag, room, header

def should_skip(slide_num, tag, room, blob):
    if slide_num <= 3:
        return "cover/legend"
    u = (tag + " " + room + " " + blob).upper()
    if re.search(r"\bHOLD\b", u):
        return "HOLD"
    if re.search(r"\b-\s*NO\b|\bNO\b.*FROST|\bNO/\s*FROST", u) and "VALANCE" not in u:
        # Slide 8 Kitchen Range NO; 17 Master Hall NO; 29 Suite 4 bath NO
        if slide_num in (8, 17, 29):
            return "NO"
    if slide_num == 8:
        return "NO"
    if slide_num == 17:
        return "NO"
    if slide_num == 29:
        return "NO"
    if not tag and not room:
        return "empty"
    return None

def build_specs(cat, blob):
    b = blob.lower()
    specs = {}
    custom = {}
    if cat == "roller":
        if "solar" in b:
            specs["rollerMaterial"] = "Solar"
        elif "woven" in b:
            specs["rollerMaterial"] = "Woven"
        elif "blackout" in b or "black out" in b:
            specs["rollerMaterial"] = "Blackout"
        else:
            specs["rollerMaterial"] = "Fabric"
        if "dual" in b:
            specs["rollerPairing"] = "Dual roller (other combo — describe in intent)"
        else:
            specs["rollerPairing"] = "Single tube"
        specs["function"] = "Operable"
        if "motor" in b or "hardwire" in b or "hard wire" in b:
            specs["operation"] = "Motorized"
        elif "manual" in b and "left" in b:
            specs["operation"] = "Continuous loop chain - LH"
        elif "manual" in b and "right" in b:
            specs["operation"] = "Continuous loop chain - RH"
        elif "manual" in b:
            specs["operation"] = "Continuous loop chain - LH"
        else:
            specs["operation"] = "Motorized" if "hard" in b else "Continuous loop chain - LH"
        specs["lining"] = "Fabric only"
        specs["mount"] = "Ceiling recess · Roll off front"
        specs["valanceStyle"] = "Custom"
        specs["stack"] = "Custom"
        custom["valanceStyle"] = "Metal fascia / housing per PPT; fascia color TBD"
        custom["stack"] = "Recessed fascia / drywall box when noted"
    elif cat == "roman":
        specs["style"] = "Flat fold"
        specs["function"] = "Operable"
        if "motor" in b or "automat" in b or "hard wire" in b or "hardwire" in b:
            specs["operation"] = "Motorized"
        elif "right" in b:
            specs["operation"] = "Continuous loop chain - RH"
        else:
            specs["operation"] = "Continuous loop chain - LH"
        specs["lining"] = "Privacy"
        specs["mount"] = "Inside · Roll off front"
        specs["valanceStyle"] = "None"
    elif cat == "drapery":
        specs["function"] = "Operable"
        specs["pleat"] = "Pinch pleat - French"
        specs["fullness"] = "2x" if "1.5" not in b and "1 ½" not in b and "1 1/2" not in b else "1.5x"
        if "1" in b and "fullness" in b:
            specs["fullness"] = "1.5x"
        specs["lining"] = "Sheer lining" if "sheer" in b else "Privacy liner"
        specs["mount"] = "Wall rod"
        specs["hem"] = "Touching"
        specs["valanceStyle"] = "Straight valance" if "valance" in b else "None"
        if "valance only" in b or "decorative valance" in b:
            specs["function"] = "Decorative non-op"
            specs["valanceStyle"] = "Straight valance"
    return specs, custom

def main():
    prs = Presentation(PPT)
    rows = []
    skipped = []
    for i, slide in enumerate(prs.slides):
        slide_num = i + 1
        texts = []
        for shape in slide.shapes:
            texts.extend(walk(shape))
        blob = "\n".join(texts)
        tag, room, header = parse_header(texts)
        # Fix combined header boxes
        for t in texts:
            if "Location:" in t and ("Work Order" in t or "ork Order" in t):
                mloc = re.search(r"Location:\s*(.+)", t, re.I)
                if mloc:
                    room = clean(mloc.group(1).split("\n")[0])
                mtag = re.search(r"(?:WINDOW|DRAPERY)\s+TREATMENT\s*([^\n]*)", t, re.I)
                if mtag:
                    tag = clean(mtag.group(1))
        reason = should_skip(slide_num, tag, room, blob)
        if reason:
            skipped.append({"slide": slide_num, "reason": reason, "tag": tag, "room": room})
            continue
        cat = infer_category(blob)
        specs, custom = build_specs(cat, blob)
        # Design intent = non-header substantive boxes
        intent_parts = []
        for t in texts:
            if re.search(r"^(Window Treatment Design|Fabrics|Design Specifications|Project:)", t, re.I):
                continue
            if re.search(r"work\s*order|ork order|^Location:", t, re.I):
                continue
            if len(t) > 8:
                intent_parts.append(t)
        intent = "\n\n".join(intent_parts)[:4000]
        notes = [
            "Imported from RH - Window Treatments Rev 2 8.4.2025.pptx — Slide " + str(slide_num) + ".",
            "Confirm fabrics, openness %, and hardware with Cindy before workroom release.",
        ]
        qty = "1"
        mq = re.search(r"QTY\s*(\d+)", tag + " " + blob, re.I)
        if mq:
            qty = mq.group(1)
        rows.append({
            "slide": slide_num,
            "category": cat,
            "location": tag or ("W-S" + str(slide_num)),
            "room": room or "TBD",
            "pieceQty": qty,
            "designIntent": intent,
            "specs": specs,
            "specCustomText": custom,
            "notes": notes,
            "fabNotes": ["Confirm COM / shade fabric before cutting."],
            "sameWindowTagMultipleTreatments": "&" in tag or " and " in tag.lower() or "," in tag,
            "pairedTreatmentNote": intent if ("&" in tag or "W5.2" in blob or "sheer" in blob.lower() and "roller" in blob.lower()) else "",
            "date": "8/4/2025",
            "revision": 2,
        })
    out = {"source": PPT, "workOrders": rows, "skipped": skipped}
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=2, ensure_ascii=False)
    print("Wrote", OUT)
    print("importable", len(rows), "skipped", len(skipped))
    for r in rows:
        print(" S%02d %s | %s | %s" % (r["slide"], r["category"], r["location"][:40], r["room"][:40]))
    for s in skipped:
        print(" SKIP S%02d %s %s" % (s["slide"], s["reason"], (s.get("tag") or "")[:30]))

if __name__ == "__main__":
    main()
