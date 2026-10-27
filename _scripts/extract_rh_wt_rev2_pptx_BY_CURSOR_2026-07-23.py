"""
Extract Cloud_RH_WT_CCH_Rev2.pptx — Arch plan marks + opening dims.
"""
import json
import re
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# OneDrive/SharePoint path often fails python-pptx PackageNotFoundError — prefer local copy.
PPT_SRC = r"C:\Users\cindy\CCH Design\Cloud - Rolling Hills - Documents\Working Documents\Layout Custom Work orders\Cloud_RH_WT_CCH_Rev2.pptx"
PPT_LOCAL = r"C:\dev\CCH-Platform-Deploy\cch-deploy\_scripts\_tmp_Cloud_RH_WT_CCH_Rev2.pptx"
OUT_RAW = r"C:\dev\CCH-Platform-Deploy\cch-deploy\_scripts\_tmp_rh_wt_rev2_raw_BY_CURSOR.json"
OUT_DUMP = r"C:\dev\CCH-Platform-Deploy\cch-deploy\_scripts\_tmp_rh_wt_rev2_dump_BY_CURSOR.txt"
OUT_STRUCT = r"C:\dev\CCH-Platform-Deploy\cch-deploy\_scripts\_tmp_rh_wt_rev2_struct_BY_CURSOR.json"


def walk(shape):
    lines = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            lines.extend(walk(s))
        return lines
    if shape.has_text_frame:
        t = shape.text_frame.text.strip()
        if t:
            lines.append(t)
    return lines


def clean(s):
    return re.sub(r"\s+", " ", (s or "").replace("\u00a0", " ")).strip()


def normalize_quotes(s):
    return (
        (s or "")
        .replace("″", '"')
        .replace("′", "'")
        .replace("\u2033", '"')
        .replace("\u2032", "'")
        .replace("“", '"')
        .replace("”", '"')
        .replace("’", "'")
    )


def feet_inches_to_inch_str(token):
    """Convert 22'-6\" or 9'-0\" or 72\" → inches with quote, e.g. 270\"."""
    if not token:
        return ""
    t = clean(normalize_quotes(token))
    m = re.match(r"^(\d+)\s*'\s*-?\s*(\d+)\s*\"?$", t)
    if m:
        return str(int(m.group(1)) * 12 + int(m.group(2))) + '"'
    m = re.match(r'^(\d+(?:\.\d+)?)\s*"?$', t)
    if m:
        return m.group(1) + '"'
    return t


def parse_arch_line(texts):
    """
    Arch: D-22 | 22'-6\" x 9'-0\" SLIDING (3 panels)
    Arch: W-L2 | 5-panel corner assy x 6'-0\" H
    Arch: (verify mark on plan)
    Arch: W-A3 + B3 (6x6) + C3/D3/E3 (2x2) small
    """
    arch_raw = ""
    for t in texts:
        for line in t.splitlines():
            m = re.match(r"^Arch:\s*(.*)$", clean(line), re.I)
            if m:
                arch_raw = clean(normalize_quotes(m.group(1)))
                break
        if arch_raw:
            break
    if not arch_raw:
        return {
            "archLine": "",
            "archMark": "",
            "openingW": "",
            "openingH": "",
            "openingNote": "",
            "verify": False,
        }

    verify = bool(re.search(r"verify|n/a|cancelled", arch_raw, re.I))
    if "|" in arch_raw:
        mark, rest = [clean(x) for x in arch_raw.split("|", 1)]
    else:
        mark, rest = arch_raw, ""

    opening_w = opening_h = ""
    opening_note = rest

    # Standard: 22'-6" x 9'-0" TYPE…
    m = re.search(
        r"(\d+\s*'\s*-?\s*\d+\s*\"?|\d+\s*\")\s*/\s*(\d+\s*'\s*-?\s*\d+\s*\"?|\d+\s*\")\s*[xX×]\s*"
        r"(\d+\s*'\s*-?\s*\d+\s*\"?|\d+\s*\")",
        rest,
    )
    if m:
        # dual widths e.g. 4'-6"/5'-0" x 6'-0"
        w1 = feet_inches_to_inch_str(m.group(1))
        w2 = feet_inches_to_inch_str(m.group(2))
        opening_h = feet_inches_to_inch_str(m.group(3))
        opening_w = w1 + " / " + w2
        opening_note = clean(rest[m.end() :]) or opening_note
    else:
        m = re.search(
            r"(\d+\s*'\s*-?\s*\d+\s*\"?|\d+\s*\")\s*[xX×]\s*(\d+\s*'\s*-?\s*\d+\s*\"?|\d+\s*\")",
            rest,
        )
        if m:
            opening_w = feet_inches_to_inch_str(m.group(1))
            opening_h = feet_inches_to_inch_str(m.group(2))
            opening_note = clean(rest[m.end() :]) or clean(rest)
        else:
            # Height-only: … x 6'-0" H
            m = re.search(r"[xX×]\s*(\d+\s*'\s*-?\s*\d+\s*\"?|\d+\s*\")\s*H\b", rest, re.I)
            if m:
                opening_h = feet_inches_to_inch_str(m.group(1))
                opening_note = clean(rest)
            else:
                # Parenthetical sizes like (6x6) — keep whole rest as note
                opening_note = rest or arch_raw

    return {
        "archLine": "Arch: " + arch_raw,
        "archMark": mark,
        "openingW": opening_w,
        "openingH": opening_h,
        "openingNote": opening_note,
        "verify": verify,
    }


def infer_category(blob):
    b = blob
    for phrase in (
        "Box to hide roller shade. Drywall",
        "Box to hide roller shade",
        "Window Treatment Design",
        "Fabrics",
        "Design Specifications",
    ):
        b = re.sub(re.escape(phrase), " ", b, flags=re.I)
    b = b.lower()
    if "valance only" in b or "decorative valance" in b or "valance over tub" in b:
        return "drapery"
    if "roman shade" in b or re.search(r"\broman\b", b):
        return "roman"
    if "draper" in b or "drapery" in b:
        if ("solar shade" in b or "roller solar" in b or "black out privacy shade" in b or "dual shade" in b) and "sheer draper" not in b:
            return "roller"
        return "drapery"
    if "dual shade" in b or "solar shade" in b or "roller solar" in b or "black out privacy shade" in b or "roller shade" in b:
        return "roller"
    return "roller"


def parse_header(texts):
    tag = ""
    room = ""
    for t in texts:
        if "Location:" in t and ("Work Order" in t or "ork Order" in t):
            mloc = re.search(r"Location:\s*(.+)", t, re.I)
            if mloc:
                room = clean(mloc.group(1).split("\n")[0])
            mtag = re.search(r"(?:WINDOW|DRAPERY)\s+TREATMENT\s*([^\n]*)", t, re.I)
            if mtag:
                tag = clean(mtag.group(1))
        elif re.search(r"^Location:", t, re.I):
            room = clean(re.sub(r"^Location:\s*", "", t, flags=re.I))
        elif re.search(r"work\s*order|ork order", t, re.I):
            m = re.search(r"(?:WINDOW|DRAPERY)\s+TREATMENT\s*(.*)$", t, re.I | re.S)
            if m:
                tag = clean(m.group(1).split("\n")[0])
    return tag, room


def should_skip(slide_num, tag, room, blob):
    if slide_num <= 3:
        return "cover/legend"
    u = (tag + " " + room + " " + blob).upper()
    if re.search(r"\bHOLD\b", u):
        return "HOLD"
    if slide_num in (8, 17, 29):
        return "NO"
    if not tag and not room:
        return "empty"
    return None


def build_measurements(cat, arch):
    meas = {}
    w = arch.get("openingW") or ""
    h = arch.get("openingH") or ""
    note_parts = []
    if arch.get("archMark"):
        note_parts.append("Plan #: " + arch["archMark"])
    if arch.get("openingNote"):
        note_parts.append(arch["openingNote"])
    if arch.get("verify"):
        note_parts.append("(verify on plan)")
    if note_parts:
        meas["NOTES"] = " | ".join(note_parts)
    if cat == "drapery":
        if w:
            meas["WINDOW W"] = w
        if h:
            meas["WINDOW H"] = h
    else:
        if w:
            meas["OPENING W"] = w
        if h:
            meas["OPENING H"] = h
    return meas


def build_specs(cat, blob):
    b = blob.lower()
    specs = {}
    custom = {}
    if cat == "roller":
        if "solar" in b:
            specs["rollerMaterial"] = "Solar"
        elif "woven" in b:
            specs["rollerMaterial"] = "Woven"
        elif "blackout" in b or "black out" in b:
            specs["rollerMaterial"] = "Blackout"
        else:
            specs["rollerMaterial"] = "Fabric"
        specs["rollerPairing"] = "Dual roller (other combo — describe in intent)" if "dual" in b else "Single tube"
        specs["function"] = "Operable"
        if "motor" in b or "hardwire" in b or "hard wire" in b:
            specs["operation"] = "Motorized"
        elif "manual" in b and "right" in b:
            specs["operation"] = "Continuous loop chain - RH"
        elif "manual" in b:
            specs["operation"] = "Continuous loop chain - LH"
        else:
            specs["operation"] = "Continuous loop chain - LH"
        specs["lining"] = "Fabric only"
        specs["mount"] = "Ceiling recess · Roll off front"
        specs["valanceStyle"] = "Custom"
        specs["stack"] = "Custom"
        custom["valanceStyle"] = "Metal fascia / housing per PPT; fascia color TBD"
        custom["stack"] = "Recessed fascia / drywall box when noted"
    elif cat == "roman":
        specs["style"] = "Flat fold"
        specs["function"] = "Operable"
        if "motor" in b or "automat" in b or "hard wire" in b or "hardwire" in b:
            specs["operation"] = "Motorized"
        elif "right" in b:
            specs["operation"] = "Continuous loop chain - RH"
        else:
            specs["operation"] = "Continuous loop chain - LH"
        specs["lining"] = "Privacy"
        specs["mount"] = "Inside · Roll off front"
        specs["valanceStyle"] = "None"
    elif cat == "drapery":
        specs["function"] = "Operable"
        specs["pleat"] = "Pinch pleat - French"
        specs["fullness"] = "1.5x" if ("1 ½" in blob or "1 1/2" in blob or "1.5" in b) else "2x"
        specs["lining"] = "Sheer lining" if "sheer" in b else "Privacy liner"
        specs["mount"] = "Wall rod"
        specs["hem"] = "Touching"
        specs["valanceStyle"] = "Straight valance" if "valance" in b else "None"
        if "valance only" in b or "decorative valance" in b:
            specs["function"] = "Decorative non-op"
            specs["valanceStyle"] = "Straight valance"
    return specs, custom


def resolve_ppt():
    import os
    import shutil
    if os.path.isfile(PPT_LOCAL):
        return PPT_LOCAL
    if os.path.isfile(PPT_SRC):
        print("Copying PPT to local path for python-pptx…")
        shutil.copy2(PPT_SRC, PPT_LOCAL)
        return PPT_LOCAL
    raise FileNotFoundError("Missing PPT at %s or %s" % (PPT_LOCAL, PPT_SRC))


def main():
    ppt = resolve_ppt()
    print("Opening", ppt)
    prs = Presentation(ppt)
    raw_slides = []
    dump_lines = []
    rows = []
    skipped = []

    for i, slide in enumerate(prs.slides):
        slide_num = i + 1
        texts = []
        for shape in slide.shapes:
            texts.extend(walk(shape))
        blob = "\n".join(texts)
        raw_slides.append({"slide": slide_num, "texts": texts})
        dump_lines.append("--- SLIDE %d ---" % slide_num)
        for t in texts:
            dump_lines.append(t)
            dump_lines.append("")

        tag, room = parse_header(texts)
        reason = should_skip(slide_num, tag, room, blob)
        if reason:
            skipped.append({"slide": slide_num, "reason": reason, "tag": tag, "room": room})
            continue

        cat = infer_category(blob)
        specs, custom = build_specs(cat, blob)
        arch = parse_arch_line(texts)
        measurements = build_measurements(cat, arch)

        intent_parts = []
        for t in texts:
            if re.search(r"^(Window Treatment Design|Fabrics|Design Specifications|Project:|Arch:)", t, re.I):
                continue
            if re.search(r"work\s*order|ork order|^Location:", t, re.I):
                continue
            if len(t) > 8:
                intent_parts.append(t)
        intent_body = "\n\n".join(intent_parts)[:3800]
        prefix = []
        if arch.get("archLine"):
            prefix.append(arch["archLine"])
        intent = ("\n".join(prefix) + ("\n\n" if prefix else "") + intent_body).strip()

        qty = "1"
        mq = re.search(r"QTY\s*(\d+)", tag + " " + blob, re.I)
        if mq:
            qty = mq.group(1)

        rows.append({
            "slide": slide_num,
            "category": cat,
            "location": tag or ("W-S" + str(slide_num)),
            "room": room or "TBD",
            "pieceQty": qty,
            "designIntent": intent,
            "specs": specs,
            "specCustomText": custom,
            "measurements": measurements,
            "arch": arch,
            "notes": [
                "Updated from Cloud_RH_WT_CCH_Rev2.pptx — Slide " + str(slide_num) + " (arch # + dims).",
                "Confirm fabrics, openness %, and hardware with Cindy before workroom release.",
            ],
            "fabNotes": ["Confirm COM / shade fabric before cutting."],
            "sameWindowTagMultipleTreatments": "&" in tag or " and " in tag.lower() or "," in tag,
            "date": "8/4/2025",
            "revision": 2,
            "sourceFile": "Cloud_RH_WT_CCH_Rev2.pptx",
        })

    with open(OUT_RAW, "w", encoding="utf-8") as f:
        json.dump({"source": ppt, "slides": raw_slides}, f, indent=2, ensure_ascii=False)
    with open(OUT_DUMP, "w", encoding="utf-8") as f:
        f.write("\n".join(dump_lines))
    with open(OUT_STRUCT, "w", encoding="utf-8") as f:
        json.dump({"source": ppt, "workOrders": rows, "skipped": skipped}, f, indent=2, ensure_ascii=False)

    print("Wrote", OUT_STRUCT)
    print("importable", len(rows), "skipped", len(skipped))
    for r in rows:
        a = r.get("arch") or {}
        print(
            " S%02d %s | %s | mark=%s | %sx%s | %s"
            % (
                r["slide"],
                r["category"],
                r["location"][:28],
                (a.get("archMark") or "-")[:28],
                a.get("openingW") or "-",
                a.get("openingH") or "-",
                (a.get("openingNote") or "")[:40],
            )
        )
    for s in skipped:
        print(" SKIP S%02d %s %s" % (s["slide"], s["reason"], (s.get("tag") or "")[:30]))


if __name__ == "__main__":
    main()
